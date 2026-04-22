from __future__ import annotations

from pathlib import Path
from shutil import copy2
from textwrap import fill

import matplotlib.pyplot as plt
import networkx as nx
import numpy as np
import pandas as pd
import seaborn as sns
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
BAB4_DIR = ROOT / "bab 4"
VIS_DIR = BAB4_DIR / "visualisasi"
DOCX_PATH = BAB4_DIR / "bab 4.docx"

FONT_NAME = "Times New Roman"

FIGURE_COPY_MAP = {
    ROOT / "data_preparation" / "visualisasi" / "02_filter_durasi.png": "07_filter_durasi_perjalanan.png",
    ROOT / "data_preparation" / "visualisasi" / "03_filter_jam_operasional.png": "08_filter_jam_operasional.png",
    ROOT / "data_preparation" / "visualisasi" / "04_feature_engineering_hari.png": "10_feature_engineering_hari.png",
    ROOT / "data_preparation" / "visualisasi" / "05_imputasi_corridor_missing.png": "11_imputasi_nilai_hilang.png",
    ROOT / "data_preparation" / "visualisasi" / "06_data_cleaning_funnel.png": "12_data_cleaning_funnel.png",
    ROOT / "data_preparation" / "visualisasi" / "07_zscore_normalisasi.png": "13_normalisasi_variabel.png",
    ROOT / "r-gmm" / "visualisasi" / "01_bic_elbow.png": "15_kurva_bic_vs_jumlah_klaster.png",
    ROOT / "r-gmm" / "visualisasi" / "02_bic_delta.png": "16_delta_bic_antar_klaster.png",
    ROOT / "r-gmm" / "visualisasi" / "step7" / "02_silhouette_score.png": "17_silhouette_score_per_jumlah_klaster.png",
    ROOT / "r-gmm" / "visualisasi" / "06_evaluation_metrics.png": "19_evaluasi_kualitas_clustering.png",
    ROOT / "r-gmm" / "visualisasi" / "03_cluster_distribution.png": "20_distribusi_ukuran_klaster.png",
    ROOT / "r-gmm" / "visualisasi" / "04_cluster_heatmap.png": "21_heatmap_karakteristik_klaster.png",
    ROOT / "r-gmm" / "visualisasi" / "05_hourly_per_cluster.png": "22_pola_jam_tap_in_per_klaster.png",
    ROOT / "r-gmm" / "visualisasi" / "09_weekend_commuter.png": "23_profil_weekend_dan_commuter.png",
    ROOT / "r-arm" / "visualisasi" / "01_rules_per_cluster.png": "25_jumlah_rules_per_klaster.png",
}

EXPECTED_VISUALS = {
    "01_top_10_missing_value.png",
    "02_top_10_koridor.png",
    "03_top_10_halte_asal.png",
    "04_top_10_halte_tujuan.png",
    "05_distribusi_jam_tap_in.png",
    "06_distribusi_durasi_perjalanan.png",
    "07_filter_durasi_perjalanan.png",
    "08_filter_jam_operasional.png",
    "09_distribusi_intensitas_perjalanan.png",
    "10_feature_engineering_hari.png",
    "11_imputasi_nilai_hilang.png",
    "12_data_cleaning_funnel.png",
    "13_normalisasi_variabel.png",
    "14_heatmap_korelasi_fitur.png",
    "15_kurva_bic_vs_jumlah_klaster.png",
    "16_delta_bic_antar_klaster.png",
    "17_silhouette_score_per_jumlah_klaster.png",
    "18_distribusi_probabilitas_posterior.png",
    "19_evaluasi_kualitas_clustering.png",
    "20_distribusi_ukuran_klaster.png",
    "21_heatmap_karakteristik_klaster.png",
    "22_pola_jam_tap_in_per_klaster.png",
    "23_profil_weekend_dan_commuter.png",
    "24_frekuensi_item_dominan_per_klaster.png",
    "25_jumlah_rules_per_klaster.png",
    "26_plot_support_confidence_lift.png",
    "27_top_od_patterns_per_klaster.png",
    "28_jaringan_aturan_asosiasi.png",
    "29_profil_temporal_klaster.png",
    "30_heatmap_dominansi_rute_per_klaster.png",
}


def fmt_int(value: int | float) -> str:
    return f"{int(round(float(value))):,}".replace(",", ".")


def fmt_num(value: int | float, digits: int = 2) -> str:
    text = f"{float(value):,.{digits}f}"
    return text.replace(",", "_").replace(".", ",").replace("_", ".")


def fmt_pct(value: int | float, digits: int = 2) -> str:
    return f"{float(value):.{digits}f}%".replace(".", ",")


def hour_label(value: int | float) -> str:
    return f"{int(round(float(value))):02d}.00"


def wrap_values(values: list[str] | pd.Index | pd.Series, width: int = 24) -> list[str]:
    return [fill(str(value), width=width) for value in list(values)]


def ensure_dirs() -> None:
    BAB4_DIR.mkdir(parents=True, exist_ok=True)
    VIS_DIR.mkdir(parents=True, exist_ok=True)


def cleanup_stale_visuals() -> None:
    for path in VIS_DIR.glob("*.png"):
        if path.name not in EXPECTED_VISUALS:
            path.unlink()


def setup_style() -> None:
    sns.set_theme(style="whitegrid", context="talk")
    plt.rcParams["figure.dpi"] = 150
    plt.rcParams["savefig.dpi"] = 200
    plt.rcParams["font.family"] = "DejaVu Sans"
    plt.rcParams["axes.titlesize"] = 18
    plt.rcParams["axes.labelsize"] = 12
    plt.rcParams["xtick.labelsize"] = 10
    plt.rcParams["ytick.labelsize"] = 10


def save(fig: plt.Figure, filename: str) -> None:
    fig.tight_layout()
    fig.savefig(VIS_DIR / filename, bbox_inches="tight", pad_inches=0.08, facecolor="white")
    plt.close(fig)


def count_csv_rows(path: Path) -> int:
    with path.open("r", encoding="utf-8", errors="ignore") as handle:
        return max(sum(1 for _ in handle) - 1, 0)


def copy_visuals() -> None:
    for source, target_name in FIGURE_COPY_MAP.items():
        copy2(source, VIS_DIR / target_name)


def load_data() -> dict[str, object]:
    raw = pd.read_csv(ROOT / "tj180.csv")
    clean = pd.read_csv(ROOT / "data_clean.csv")
    clean["date"] = pd.to_datetime(clean["date"])
    clean["week_part"] = np.where(clean["is_weekend"] == 1, "Akhir Pekan", "Hari Kerja")

    step00 = pd.read_csv(ROOT / "data_preparation" / "csv_outputs" / "STEP_00_data_quality.csv")
    step04 = pd.read_csv(ROOT / "data_preparation" / "csv_outputs" / "STEP_04_features.csv")
    step05 = pd.read_csv(ROOT / "data_preparation" / "csv_outputs" / "STEP_05_imputed.csv")
    step06 = pd.read_csv(ROOT / "data_preparation" / "csv_outputs" / "STEP_06_cleaned.csv")
    step07 = pd.read_csv(ROOT / "data_preparation" / "csv_outputs" / "STEP_07_normalized.csv")

    cluster_labeled = pd.read_csv(ROOT / "r-gmm" / "hasil" / "06_cluster_labeled.csv")
    cluster_probabilities = pd.read_csv(ROOT / "r-gmm" / "hasil" / "05_cluster_probabilities.csv")
    cluster_profiles = pd.read_csv(ROOT / "r-gmm" / "hasil" / "06_cluster_profiles.csv")
    model_selection = pd.read_csv(ROOT / "r-gmm" / "hasil" / "04_model_selection.csv")
    bic_best = pd.read_csv(ROOT / "r-gmm" / "hasil" / "04_bic_best_per_k.csv")
    eval_scores = pd.read_csv(ROOT / "r-gmm" / "hasil" / "07_evaluation_scores.csv")

    item_freq_cluster = pd.read_csv(ROOT / "r-arm" / "hasil" / "02_item_frequency_cluster.csv")
    trip_cluster_base = pd.read_csv(ROOT / "r-arm" / "hasil" / "01_trip_cluster_base.csv")
    rules_cluster = pd.read_csv(ROOT / "r-arm" / "hasil" / "04_rules_cluster_filtered.csv")
    arm_eval = pd.read_csv(ROOT / "r-arm" / "hasil" / "06_arm_evaluation.csv")
    arm_eval_cluster = pd.read_csv(ROOT / "r-arm" / "hasil" / "06_arm_evaluation_cluster.csv")
    filter_params = pd.read_csv(ROOT / "r-arm" / "hasil" / "04_filter_params.csv")

    stage_counts = {
        "Dataset awal": count_csv_rows(ROOT / "tj180.csv"),
        "Setelah parsing waktu": count_csv_rows(ROOT / "data_preparation" / "csv_outputs" / "STEP_01_parsed.csv"),
        "Setelah filter durasi": count_csv_rows(ROOT / "data_preparation" / "csv_outputs" / "STEP_02_filter_durasi.csv"),
        "Setelah filter jam": count_csv_rows(ROOT / "data_preparation" / "csv_outputs" / "STEP_03_filter_jam.csv"),
        "Setelah feature engineering": count_csv_rows(ROOT / "data_preparation" / "csv_outputs" / "STEP_04_features.csv"),
        "Setelah imputasi": count_csv_rows(ROOT / "data_preparation" / "csv_outputs" / "STEP_05_imputed.csv"),
        "Setelah cleaning": count_csv_rows(ROOT / "data_preparation" / "csv_outputs" / "STEP_06_cleaned.csv"),
        "Setelah normalisasi": count_csv_rows(ROOT / "data_preparation" / "csv_outputs" / "STEP_07_normalized.csv"),
        "Dataset final": count_csv_rows(ROOT / "data_clean.csv"),
    }

    return {
        "raw": raw,
        "clean": clean,
        "step00": step00,
        "step04": step04,
        "step05": step05,
        "step06": step06,
        "step07": step07,
        "cluster_labeled": cluster_labeled,
        "cluster_probabilities": cluster_probabilities,
        "cluster_profiles": cluster_profiles,
        "model_selection": model_selection,
        "bic_best": bic_best,
        "eval_scores": eval_scores,
        "item_freq_cluster": item_freq_cluster,
        "trip_cluster_base": trip_cluster_base,
        "rules_cluster": rules_cluster,
        "arm_eval": arm_eval,
        "arm_eval_cluster": arm_eval_cluster,
        "filter_params": filter_params,
        "stage_counts": stage_counts,
    }


def plot_top10_missing(step00: pd.DataFrame) -> None:
    plot_df = step00.sort_values("missing_values", ascending=False).head(10).copy()
    fig, ax = plt.subplots(figsize=(12, 7))
    sns.barplot(
        data=plot_df,
        x="missing_values",
        y="kolom",
        hue="kolom",
        palette="rocket",
        dodge=False,
        legend=False,
        ax=ax,
    )
    ax.set_title("Top 10 Missing Value pada Dataset Awal")
    ax.set_xlabel("Jumlah Missing Value")
    ax.set_ylabel("")
    ax.invert_yaxis()
    for i, row in plot_df.reset_index(drop=True).iterrows():
        ax.text(float(row["missing_values"]) + 150, i, f"{fmt_int(row['missing_values'])} ({fmt_pct(row['pct'])})", va="center", fontsize=10)
    save(fig, "01_top_10_missing_value.png")


def plot_top10_bar(series: pd.Series, title: str, xlabel: str, filename: str, palette: str) -> None:
    top = series.value_counts().head(10)
    labels = wrap_values(top.index, 26)
    fig, ax = plt.subplots(figsize=(12, 7))
    sns.barplot(
        x=top.values,
        y=labels,
        hue=labels,
        palette=palette,
        dodge=False,
        legend=False,
        ax=ax,
    )
    ax.set_title(title)
    ax.set_xlabel(xlabel)
    ax.set_ylabel("")
    ax.invert_yaxis()
    for i, value in enumerate(top.values):
        ax.text(float(value) + max(top.values) * 0.01, i, fmt_int(value), va="center", fontsize=10)
    save(fig, filename)


def plot_tapin_hist(clean: pd.DataFrame) -> None:
    bins = np.arange(5, 23.5, 0.5)
    fig, ax = plt.subplots(figsize=(12, 7))
    palette = {"Hari Kerja": "#0f766e", "Akhir Pekan": "#b45309"}
    for label, color in palette.items():
        subset = clean.loc[clean["week_part"] == label, "tapIn_hour"]
        ax.hist(subset, bins=bins, density=True, alpha=0.55, color=color, label=label)
    ax.set_title("Distribusi Jam Tap-In pada Hari Kerja dan Akhir Pekan")
    ax.set_xlabel("Jam Tap-In")
    ax.set_ylabel("Kepadatan")
    ax.legend()
    save(fig, "05_distribusi_jam_tap_in.png")


def plot_duration_hist(clean: pd.DataFrame) -> None:
    fig, ax = plt.subplots(figsize=(12, 7))
    sns.histplot(clean["duration_minutes"], bins=45, kde=True, color="#2563eb", ax=ax)
    ax.set_title("Distribusi Durasi Perjalanan")
    ax.set_xlabel("Durasi (menit)")
    ax.set_ylabel("Frekuensi")
    save(fig, "06_distribusi_durasi_perjalanan.png")


def plot_intensity_hist(clean: pd.DataFrame) -> None:
    intensity = clean[["payCardID", "date", "n_trips"]].drop_duplicates()
    bins = np.arange(0.5, intensity["n_trips"].max() + 1.5, 1)
    fig, ax = plt.subplots(figsize=(12, 7))
    sns.histplot(intensity["n_trips"], bins=bins, discrete=True, color="#7c3aed", ax=ax)
    ax.set_title("Distribusi Intensitas Perjalanan per Pengguna per Hari")
    ax.set_xlabel("Jumlah Trip per Pengguna per Hari")
    ax.set_ylabel("Frekuensi")
    save(fig, "09_distribusi_intensitas_perjalanan.png")


def plot_correlation_heatmap(clean: pd.DataFrame) -> None:
    corr_cols = [
        "tapIn_hour",
        "tapOut_hour",
        "duration_minutes",
        "day_of_week",
        "is_weekend",
        "n_trips",
        "n_days_month",
        "is_commuter",
        "trip_num",
    ]
    corr = clean[corr_cols].corr(numeric_only=True)
    fig, ax = plt.subplots(figsize=(11, 8))
    sns.heatmap(corr, annot=True, fmt=".2f", cmap="YlGnBu", square=True, annot_kws={"size": 7}, ax=ax)
    ax.set_title("Heatmap Korelasi Antar Fitur")
    save(fig, "14_heatmap_korelasi_fitur.png")


def plot_posterior_distribution(cluster_probabilities: pd.DataFrame) -> None:
    prob_cols = [col for col in cluster_probabilities.columns if col.startswith("prob_cl")]
    probs = cluster_probabilities[prob_cols].astype(float)
    max_prob = probs.max(axis=1)
    if max_prob.max() > 1:
        max_prob = max_prob / 100.0
    fig, ax = plt.subplots(figsize=(12, 7))
    sns.histplot(max_prob, bins=30, kde=True, color="#0f766e", ax=ax)
    ax.axvline(max_prob.median(), color="#991b1b", linestyle="--", linewidth=2)
    ax.axvline(0.9, color="#1d4ed8", linestyle=":", linewidth=2)
    ax.set_title("Distribusi Probabilitas Posterior Maksimum")
    ax.set_xlabel("Probabilitas Maksimum")
    ax.set_ylabel("Frekuensi")
    save(fig, "18_distribusi_probabilitas_posterior.png")


def plot_item_frequency_cluster(item_freq_cluster: pd.DataFrame) -> None:
    top_items = (
        item_freq_cluster.sort_values(["cluster", "support"], ascending=[True, False])
        .groupby("cluster_label")
        .head(3)
        .copy()
    )
    top_items["item_wrap"] = wrap_values(top_items["item"], 24)
    clusters = list(top_items["cluster_label"].unique())
    fig, axes = plt.subplots(len(clusters), 1, figsize=(12, 16), sharex=False)
    if len(clusters) == 1:
        axes = [axes]
    for ax, cluster in zip(axes, clusters):
        subset = top_items[top_items["cluster_label"] == cluster]
        sns.barplot(
            data=subset,
            x="support",
            y="item_wrap",
            hue="item_wrap",
            palette="flare",
            dodge=False,
            legend=False,
            ax=ax,
        )
        ax.set_title(cluster, loc="left", fontsize=13, weight="bold")
        ax.set_xlabel("Support")
        ax.set_ylabel("")
        ax.invert_yaxis()
        for i, value in enumerate(subset["support"]):
            ax.text(float(value) + subset["support"].max() * 0.02, i, fmt_pct(value * 100), va="center", fontsize=9)
    fig.suptitle("Frekuensi Item Dominan per Klaster", y=1.01, fontsize=18, weight="bold")
    save(fig, "24_frekuensi_item_dominan_per_klaster.png")


def plot_support_confidence_lift(rules_cluster: pd.DataFrame) -> None:
    plot_df = rules_cluster.copy()
    plot_df["support_local_pct"] = plot_df["support_local"].astype(float) * 100
    plot_df["confidence_pct"] = plot_df["confidence"].astype(float) * 100
    plot_df["lift_local"] = plot_df["lift_local"].astype(float)
    fig, ax = plt.subplots(figsize=(12, 8))
    sns.scatterplot(
        data=plot_df,
        x="support_local_pct",
        y="confidence_pct",
        size="lift_local",
        hue="cluster_label",
        sizes=(90, 900),
        alpha=0.65,
        palette="tab10",
        ax=ax,
    )
    ax.set_title("Plot Support-Confidence-Lift")
    ax.set_xlabel("Support Lokal (%)")
    ax.set_ylabel("Confidence (%)")
    save(fig, "26_plot_support_confidence_lift.png")


def plot_top_od_patterns(trip_cluster_base: pd.DataFrame) -> None:
    od = (
        trip_cluster_base.groupby(["cluster_label", "lhs", "rhs"])
        .size()
        .reset_index(name="count")
        .sort_values(["cluster_label", "count"], ascending=[True, False])
        .groupby("cluster_label")
        .head(3)
        .copy()
    )
    od["od_wrap"] = wrap_values(od["lhs"] + " -> " + od["rhs"], 30)
    clusters = list(od["cluster_label"].unique())
    fig, axes = plt.subplots(len(clusters), 1, figsize=(12, 18), sharex=False)
    if len(clusters) == 1:
        axes = [axes]
    for ax, cluster in zip(axes, clusters):
        subset = od[od["cluster_label"] == cluster]
        sns.barplot(
            data=subset,
            x="count",
            y="od_wrap",
            hue="od_wrap",
            palette="crest",
            dodge=False,
            legend=False,
            ax=ax,
        )
        ax.set_title(cluster, loc="left", fontsize=13, weight="bold")
        ax.set_xlabel("Jumlah Perjalanan")
        ax.set_ylabel("")
        ax.invert_yaxis()
        for i, value in enumerate(subset["count"]):
            ax.text(float(value) + subset["count"].max() * 0.02, i, fmt_int(value), va="center", fontsize=9)
    fig.suptitle("Top Origin-Destination Patterns per Klaster", y=1.01, fontsize=18, weight="bold")
    save(fig, "27_top_od_patterns_per_klaster.png")


def plot_rule_network(rules_cluster: pd.DataFrame) -> None:
    filtered = rules_cluster.copy()
    filtered["lift_local"] = filtered["lift_local"].astype(float)
    edges = filtered[(filtered["lift_local"] > 1) & (filtered["lhs"] != filtered["rhs"])].sort_values("rule_score", ascending=False).head(24)
    if edges.empty:
        edges = filtered.sort_values("rule_score", ascending=False).head(24)

    graph = nx.Graph()
    palette = dict(zip(edges["cluster_label"].unique(), sns.color_palette("tab10", n_colors=edges["cluster_label"].nunique())))
    for _, row in edges.iterrows():
        lhs = str(row["lhs"])
        rhs = str(row["rhs"])
        graph.add_node(lhs)
        graph.add_node(rhs)
        graph.add_edge(lhs, rhs, weight=float(row["lift_local"]), cluster=str(row["cluster_label"]))

    fig, ax = plt.subplots(figsize=(13, 10))
    ax.set_title("Jaringan Aturan Asosiasi", fontsize=18, weight="bold")
    ax.axis("off")
    if graph.number_of_nodes() == 0:
        ax.text(0.5, 0.5, "Tidak ada aturan yang dapat divisualisasikan", ha="center", va="center")
        save(fig, "28_jaringan_aturan_asosiasi.png")
        return

    pos = nx.spring_layout(graph, k=1.0, seed=42)
    degrees = dict(graph.degree())
    node_sizes = [420 + degrees[node] * 170 for node in graph.nodes()]
    nx.draw_networkx_nodes(graph, pos, node_size=node_sizes, node_color="#dbeafe", edgecolors="#1e3a8a", ax=ax)
    for cluster_name, color in palette.items():
        edge_list = [(u, v) for u, v, d in graph.edges(data=True) if d["cluster"] == cluster_name]
        widths = [graph[u][v]["weight"] * 0.18 for u, v in edge_list]
        nx.draw_networkx_edges(graph, pos, edgelist=edge_list, width=widths, edge_color=[color], alpha=0.85, ax=ax)
    labels = {node: fill(node, 16) for node in graph.nodes()}
    nx.draw_networkx_labels(graph, pos, labels=labels, font_size=8, ax=ax)
    save(fig, "28_jaringan_aturan_asosiasi.png")


def plot_temporal_profile(cluster_labeled: pd.DataFrame) -> None:
    df = cluster_labeled.copy()
    df["kategori_jam"] = pd.cut(
        df["tapIn_hour"],
        bins=[0, 10, 15, 19, 24],
        labels=["Pagi", "Siang", "Sore", "Malam"],
        include_lowest=True,
        right=False,
    )
    profile = (
        df.groupby(["kategori_jam", "label"], observed=False)
        .size()
        .reset_index(name="count")
        .pivot(index="kategori_jam", columns="label", values="count")
        .fillna(0)
    )
    profile = profile.div(profile.sum(axis=1), axis=0)
    fig, ax = plt.subplots(figsize=(12, 7))
    profile.plot(kind="bar", stacked=True, colormap="tab20", ax=ax)
    ax.set_title("Profil Temporal Klaster")
    ax.set_xlabel("Kategori Waktu")
    ax.set_ylabel("Proporsi")
    ax.legend(title="Klaster", bbox_to_anchor=(1.02, 1), loc="upper left")
    save(fig, "29_profil_temporal_klaster.png")


def plot_route_heatmap(trip_cluster_base: pd.DataFrame) -> None:
    route_counts = trip_cluster_base.groupby(["cluster_label", "lhs"]).size().reset_index(name="count")
    top_routes = route_counts.groupby("lhs")["count"].sum().sort_values(ascending=False).head(12).index
    heatmap_df = route_counts[route_counts["lhs"].isin(top_routes)].pivot(index="cluster_label", columns="lhs", values="count").fillna(0)
    heatmap_df = heatmap_df.div(heatmap_df.sum(axis=1), axis=0)
    fig, ax = plt.subplots(figsize=(14, 7))
    sns.heatmap(heatmap_df, cmap="rocket_r", annot=True, fmt=".2f", annot_kws={"size": 7}, ax=ax)
    ax.set_title("Heatmap Dominansi Rute per Klaster")
    ax.set_xlabel("Rute Dominan")
    ax.set_ylabel("Klaster")
    ax.set_xticklabels([fill(str(label.get_text()), 18) for label in ax.get_xticklabels()], rotation=35, ha="right")
    save(fig, "30_heatmap_dominansi_rute_per_klaster.png")


def build_visuals(loaded: dict[str, object]) -> None:
    plot_top10_missing(loaded["step00"])
    plot_top10_bar(loaded["clean"]["corridorName"], "Top 10 Koridor dengan Frekuensi Tertinggi", "Jumlah Perjalanan", "02_top_10_koridor.png", "crest")
    plot_top10_bar(loaded["clean"]["tapInStopsName"], "Top 10 Halte Asal dengan Frekuensi Tap-In Tertinggi", "Jumlah Tap-In", "03_top_10_halte_asal.png", "mako")
    plot_top10_bar(loaded["clean"]["tapOutStopsName"], "Top 10 Halte Tujuan dengan Frekuensi Tap-Out Tertinggi", "Jumlah Tap-Out", "04_top_10_halte_tujuan.png", "flare")
    plot_tapin_hist(loaded["clean"])
    plot_duration_hist(loaded["clean"])
    plot_intensity_hist(loaded["clean"])
    plot_correlation_heatmap(loaded["clean"])
    plot_posterior_distribution(loaded["cluster_probabilities"])
    plot_item_frequency_cluster(loaded["item_freq_cluster"])
    plot_support_confidence_lift(loaded["rules_cluster"])
    plot_top_od_patterns(loaded["trip_cluster_base"])
    plot_rule_network(loaded["rules_cluster"])
    plot_temporal_profile(loaded["cluster_labeled"])
    plot_route_heatmap(loaded["trip_cluster_base"])


RAW_COLUMN_INFO = [
    ("transID", "ID unik transaksi"),
    ("payCardID", "ID kartu pembayaran anonim"),
    ("corridorID", "Kode koridor awal perjalanan"),
    ("corridorName", "Nama koridor awal perjalanan"),
    ("tapInStopsName", "Nama halte asal"),
    ("tapOutStopsName", "Nama halte tujuan"),
    ("direction", "Arah perjalanan"),
    ("tapInTime", "Waktu tap-in mentah"),
    ("tapOutTime", "Waktu tap-out mentah"),
    ("payAmount", "Nilai pembayaran transaksi"),
]

GMM_FEATURES = [
    ("z_tapIn_hour", "Kontinu", "Jam tap-in yang telah dinormalisasi dengan z-score"),
    ("z_duration_minutes", "Kontinu", "Durasi perjalanan yang telah dinormalisasi"),
    ("z_n_trips", "Kontinu", "Jumlah trip per pengguna per hari yang telah dinormalisasi"),
    ("z_n_days_month", "Kontinu", "Jumlah hari aktif per bulan yang telah dinormalisasi"),
    ("is_weekend", "Biner", "Penanda transaksi akhir pekan"),
    ("is_commuter", "Biner", "Penanda pengguna rutin atau commuter"),
]


def compute_stats(loaded: dict[str, object]) -> dict[str, object]:
    raw: pd.DataFrame = loaded["raw"]
    clean: pd.DataFrame = loaded["clean"]
    step00: pd.DataFrame = loaded["step00"]
    step04: pd.DataFrame = loaded["step04"]
    step05: pd.DataFrame = loaded["step05"]
    step06: pd.DataFrame = loaded["step06"]
    cluster_labeled: pd.DataFrame = loaded["cluster_labeled"]
    cluster_probabilities: pd.DataFrame = loaded["cluster_probabilities"]
    cluster_profiles: pd.DataFrame = loaded["cluster_profiles"]
    model_selection: pd.DataFrame = loaded["model_selection"]
    bic_best: pd.DataFrame = loaded["bic_best"]
    eval_scores: pd.DataFrame = loaded["eval_scores"]
    item_freq_cluster: pd.DataFrame = loaded["item_freq_cluster"]
    trip_cluster_base: pd.DataFrame = loaded["trip_cluster_base"]
    rules_cluster: pd.DataFrame = loaded["rules_cluster"]
    arm_eval: pd.DataFrame = loaded["arm_eval"]
    arm_eval_cluster: pd.DataFrame = loaded["arm_eval_cluster"]
    filter_params: pd.DataFrame = loaded["filter_params"]
    stage_counts: dict[str, int] = loaded["stage_counts"]

    stats: dict[str, object] = {}
    stats["raw_rows"] = len(raw)
    stats["raw_cols"] = raw.shape[1]
    stats["clean_rows"] = len(clean)
    stats["clean_cols"] = clean.shape[1]
    stats["period"] = f"{clean['date'].min():%d-%m-%Y} sampai {clean['date'].max():%d-%m-%Y}"
    stats["stage_counts"] = stage_counts
    stats["removed_duration"] = stage_counts["Setelah parsing waktu"] - stage_counts["Setelah filter durasi"]
    stats["removed_hour"] = stage_counts["Setelah filter durasi"] - stage_counts["Setelah filter jam"]
    stats["removed_cleaning"] = stage_counts["Setelah imputasi"] - stage_counts["Setelah cleaning"]
    stats["removed_total"] = stage_counts["Dataset awal"] - stage_counts["Dataset final"]
    stats["removed_total_pct"] = stats["removed_total"] / stage_counts["Dataset awal"] * 100

    raw_columns_table = []
    for column, desc in RAW_COLUMN_INFO:
        raw_columns_table.append(
            [
                column,
                str(raw[column].dtype),
                fmt_int(raw[column].notna().sum()),
                fmt_int(raw[column].isna().sum()),
                desc,
            ]
        )
    stats["raw_columns_table"] = raw_columns_table
    stats["gmm_feature_table"] = [[feature, kind, desc] for feature, kind, desc in GMM_FEATURES]

    top_missing = step00.sort_values("missing_values", ascending=False).head(10).copy()
    stats["top_missing"] = top_missing
    stats["top_corridors"] = clean["corridorName"].value_counts().head(10)
    stats["top_tapin"] = clean["tapInStopsName"].value_counts().head(10)
    stats["top_tapout"] = clean["tapOutStopsName"].value_counts().head(10)

    stats["peak_hours"] = clean["tapIn_hour"].round().value_counts().sort_values(ascending=False).head(8)
    stats["weekday_peak"] = clean.loc[clean["is_weekend"] == 0, "tapIn_hour"].round().value_counts().sort_values(ascending=False).head(5)
    stats["weekend_peak"] = clean.loc[clean["is_weekend"] == 1, "tapIn_hour"].round().value_counts().sort_values(ascending=False).head(5)

    duration_counts = pd.cut(clean["duration_minutes"], bins=[0, 30, 60, 90, 120, 180], right=False).value_counts().sort_index()
    duration_pct = duration_counts / len(clean) * 100
    stats["duration_table"] = list(zip(duration_counts.index.astype(str), duration_counts.tolist(), duration_pct.tolist()))

    intensity = clean[["payCardID", "date", "n_trips"]].drop_duplicates()
    intensity_counts = intensity["n_trips"].value_counts().sort_index()
    intensity_pct = intensity_counts / len(intensity) * 100
    stats["intensity_table"] = list(zip(intensity_counts.index.tolist(), intensity_counts.tolist(), intensity_pct.tolist()))

    stats["weekend_pct"] = float(clean["is_weekend"].mean() * 100)
    stats["weekday_pct"] = 100.0 - stats["weekend_pct"]
    stats["commuter_pct"] = float(clean["is_commuter"].mean() * 100)

    stats["corridor_missing_before"] = int(step04["corridorName"].isna().sum())
    stats["corridor_missing_after"] = int(step05["corridorName"].isna().sum())
    stats["corridor_missing_final"] = int(step06["corridorName"].isna().sum())
    stats["tapout_missing_after"] = int(step05["tapOutStopsName"].isna().sum())
    stats["imputation_reduction_pct"] = (stats["corridor_missing_before"] - stats["corridor_missing_after"]) / stats["corridor_missing_before"] * 100

    corr_cols = [
        "tapIn_hour",
        "tapOut_hour",
        "duration_minutes",
        "day_of_week",
        "is_weekend",
        "n_trips",
        "n_days_month",
        "is_commuter",
        "trip_num",
    ]
    corr = clean[corr_cols].corr(numeric_only=True)
    corr_pairs: list[tuple[str, str, float]] = []
    for i, col1 in enumerate(corr_cols):
        for col2 in corr_cols[i + 1 :]:
            corr_pairs.append((col1, col2, float(corr.loc[col1, col2])))
    stats["corr_pairs"] = sorted(corr_pairs, key=lambda item: abs(item[2]), reverse=True)
    stats["corr_tap_duration"] = float(corr.loc["tapIn_hour", "duration_minutes"])

    stats["selected_k"] = int(model_selection.loc[0, "selected_k"])
    stats["selected_model"] = str(model_selection.loc[0, "selected_model"])
    stats["selected_bic"] = float(model_selection.loc[0, "selected_bic"])
    bic_best = bic_best.sort_values("K").copy()
    stats["bic_map"] = dict(zip(bic_best["K"], bic_best["BIC"]))
    stats["bic_delta_45"] = float(stats["bic_map"][5] - stats["bic_map"][4])
    stats["bic_delta_56"] = float(stats["bic_map"][6] - stats["bic_map"][5])

    eval_map = eval_scores.set_index("K")
    stats["silhouette_k4"] = float(eval_map.loc[4, "Silhouette"])
    stats["silhouette_k5"] = float(eval_map.loc[5, "Silhouette"])
    stats["silhouette_k6"] = float(eval_map.loc[6, "Silhouette"])
    stats["entropy_k5"] = float(eval_map.loc[5, "Entropy"])
    stats["balance_k5"] = float(eval_map.loc[5, "Cluster_balance"])
    stats["composite_k5"] = float(eval_map.loc[5, "Composite_score"])

    prob_cols = [col for col in cluster_probabilities.columns if col.startswith("prob_cl")]
    max_prob = cluster_probabilities[prob_cols].astype(float).max(axis=1)
    if max_prob.max() > 1:
        max_prob = max_prob / 100.0
    stats["posterior_median"] = float(max_prob.median())
    stats["posterior_p75"] = float(max_prob.quantile(0.75))
    stats["posterior_ge_90"] = float((max_prob >= 0.9).mean() * 100)
    stats["posterior_ge_80"] = float((max_prob >= 0.8).mean() * 100)

    cluster_profiles = cluster_profiles.copy()
    stats["cluster_profiles"] = cluster_profiles.set_index("label")
    stats["largest_cluster"] = cluster_profiles.sort_values("n_obs", ascending=False).iloc[0]
    stats["smallest_cluster"] = cluster_profiles.sort_values("n_obs", ascending=True).iloc[0]

    arm_map = dict(zip(arm_eval["metric"], arm_eval["value"].astype(float)))
    stats["connected_trip_total"] = int(arm_map["connected_trip_total"])
    stats["pair_candidate_total"] = int(arm_map["pair_candidate_total"])
    stats["rules_selected_global"] = int(arm_map["rules_selected_global"])
    stats["trip_coverage_pct"] = float(arm_map["trip_coverage_by_rules_pct"])
    stats["avg_support_global_pct"] = float(arm_map["avg_support_global_pct"])
    stats["avg_support_local_pct"] = float(arm_map["avg_support_local_pct"])
    stats["avg_lift_local"] = float(arm_map["avg_lift_local"])
    stats["median_lift_local"] = float(arm_map["median_lift_local"])
    stats["avg_lift_global"] = float(arm_map["avg_lift_global"])
    stats["connected_pct_of_final"] = stats["connected_trip_total"] / len(clean) * 100

    rules_by_cluster = arm_eval_cluster.set_index("cluster_label")["rules_selected"]
    stats["rules_by_cluster"] = rules_by_cluster
    stats["arm_eval_cluster"] = arm_eval_cluster.set_index("cluster_label")
    stats["filter_params"] = dict(zip(filter_params["parameter"], filter_params["value"]))

    top_item_cluster = (
        item_freq_cluster.sort_values(["cluster", "support"], ascending=[True, False])
        .groupby("cluster_label")
        .head(1)
        .set_index("cluster_label")
    )
    stats["top_item_cluster"] = top_item_cluster

    top_od_cluster = (
        trip_cluster_base.groupby(["cluster_label", "lhs", "rhs"]).size().reset_index(name="count")
        .sort_values(["cluster_label", "count"], ascending=[True, False])
        .groupby("cluster_label")
        .head(1)
        .set_index("cluster_label")
    )
    stats["top_od_cluster"] = top_od_cluster

    top_rule = rules_cluster.copy().sort_values("lift_local", ascending=False).iloc[0]
    stats["top_rule"] = top_rule
    stats["cross_pct"] = float(trip_cluster_base["is_cross"].mean() * 100)
    stats["self_pct"] = 100.0 - stats["cross_pct"]

    temporal = cluster_labeled.copy()
    temporal["kategori_jam"] = pd.cut(
        temporal["tapIn_hour"],
        bins=[0, 10, 15, 19, 24],
        labels=["Pagi", "Siang", "Sore", "Malam"],
        include_lowest=True,
        right=False,
    )
    temporal_dom = temporal.groupby(["kategori_jam", "label"], observed=False).size().reset_index(name="count")
    temporal_dom["pct"] = temporal_dom.groupby("kategori_jam", observed=False)["count"].transform(lambda s: s / s.sum() * 100)
    stats["temporal_dominance"] = temporal_dom.sort_values(["kategori_jam", "pct"], ascending=[True, False])

    route_dom = trip_cluster_base.groupby(["cluster_label", "lhs"]).size().reset_index(name="count")
    route_dom["pct"] = route_dom.groupby("cluster_label")["count"].transform(lambda s: s / s.sum() * 100)
    stats["route_dominance"] = route_dom.sort_values(["cluster_label", "count"], ascending=[True, False])

    return stats


def set_run_font(run, bold: bool = False, size: int = 12) -> None:
    run.bold = bold
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)


def add_text(doc: Document, text: str = "", bold: bool = False, center: bool = False, indent: bool = True) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.first_line_indent = Cm(0) if not indent else Cm(1.25)
    para.paragraph_format.line_spacing = 1.5
    run = para.add_run(text)
    set_run_font(run, bold=bold)


def set_cell(cell, text: str, bold: bool = False, align: int = WD_ALIGN_PARAGRAPH.CENTER) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    para = cell.paragraphs[0]
    para.alignment = align
    para.paragraph_format.first_line_indent = Cm(0)
    run = para.add_run(str(text))
    set_run_font(run, bold=bold, size=11)


def add_doc_table(doc: Document, table_no: int, title: str, headers: list[str], rows: list[list[str]]) -> None:
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.line_spacing = 1.5
    run = cap.add_run(f"Tabel 4.{table_no} {title}")
    set_run_font(run)

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, header in enumerate(headers):
        set_cell(table.rows[0].cells[idx], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 or len(str(value)) > 28 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell(cells[idx], value, align=align)


def add_figure(doc: Document, fig_no: int, filename: str, caption_text: str, description: str) -> None:
    frame = doc.add_table(rows=1, cols=1)
    frame.style = "Table Grid"
    frame.alignment = WD_TABLE_ALIGNMENT.CENTER
    frame.autofit = False
    cell = frame.cell(0, 0)
    cell.width = Inches(6.45)
    p_img = cell.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.first_line_indent = Cm(0)
    p_img.paragraph_format.line_spacing = 1.0
    p_img.add_run().add_picture(str(VIS_DIR / filename), width=Inches(6.08))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.line_spacing = 1.5
    run = cap.add_run(f"Gambar 4.{fig_no} {caption_text}")
    set_run_font(run)

    add_text(doc, f"Berdasarkan Gambar 4.{fig_no}, {description}")


def build_docx(stats: dict[str, object]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(3)
    section.left_margin = Cm(4)
    section.right_margin = Cm(3)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(12)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5

    stage_counts: dict[str, int] = stats["stage_counts"]
    top_missing: pd.DataFrame = stats["top_missing"]
    top_corridors: pd.Series = stats["top_corridors"]
    top_tapin: pd.Series = stats["top_tapin"]
    top_tapout: pd.Series = stats["top_tapout"]
    peak_hours: pd.Series = stats["peak_hours"]
    weekday_peak: pd.Series = stats["weekday_peak"]
    weekend_peak: pd.Series = stats["weekend_peak"]
    duration_table: list[tuple[str, int, float]] = stats["duration_table"]
    intensity_table: list[tuple[int, int, float]] = stats["intensity_table"]
    corr_pairs: list[tuple[str, str, float]] = stats["corr_pairs"]
    cluster_profiles: pd.DataFrame = stats["cluster_profiles"]
    rules_by_cluster: pd.Series = stats["rules_by_cluster"]
    top_item_cluster: pd.DataFrame = stats["top_item_cluster"]
    top_od_cluster: pd.DataFrame = stats["top_od_cluster"]
    top_rule: pd.Series = stats["top_rule"]
    temporal_dominance: pd.DataFrame = stats["temporal_dominance"]
    route_dominance: pd.DataFrame = stats["route_dominance"]
    arm_eval_cluster: pd.DataFrame = stats["arm_eval_cluster"]

    pagi = temporal_dominance[temporal_dominance["kategori_jam"] == "Pagi"].iloc[0]
    siang = temporal_dominance[temporal_dominance["kategori_jam"] == "Siang"].iloc[0]
    sore = temporal_dominance[temporal_dominance["kategori_jam"] == "Sore"].iloc[0]
    malam = temporal_dominance[temporal_dominance["kategori_jam"] == "Malam"].iloc[0]

    add_text(doc, "BAB IV", bold=True, center=True, indent=False)
    add_text(doc, "HASIL DAN PEMBAHASAN", bold=True, center=True, indent=False)
    add_text(
        doc,
        "Bab ini menyajikan hasil penelitian sesuai urutan metodologi pada Bab III. Pembahasan dimulai dari hasil data preparation, dilanjutkan dengan hasil pemodelan Gaussian Mixture Model, hasil Association Rule Mining, dan diakhiri dengan analisis integratif terhadap pola mobilitas penumpang TransJakarta.",
    )

    add_doc_table(
        doc,
        1,
        "Sepuluh Atribut Utama pada Dataset Awal",
        ["Kolom", "Tipe Data", "Nilai Valid", "Missing", "Keterangan"],
        stats["raw_columns_table"],
    )
    add_doc_table(
        doc,
        2,
        "Perubahan Jumlah Data pada Setiap Tahap Data Preparation",
        ["Tahap", "Jumlah Baris"],
        [[stage, fmt_int(count)] for stage, count in stage_counts.items()],
    )

    add_text(doc, "4.1 Hasil Data Preparation", bold=True, indent=False)
    add_text(
        doc,
        f"Tahap data preparation menghasilkan dataset final sebanyak {fmt_int(stats['clean_rows'])} baris dari dataset awal {fmt_int(stats['raw_rows'])} baris. Secara total terdapat pengurangan {fmt_int(stats['removed_total'])} baris atau {fmt_pct(stats['removed_total_pct'])}. Dataset akhir digunakan untuk pemodelan karena sudah melalui transformasi waktu, filtering, imputasi, cleaning, dan normalisasi.",
    )

    add_text(doc, "4.1.1 Data Quality Assessment", bold=True, indent=False)
    add_text(
        doc,
        f"Dataset awal berasal dari transaksi tap-in dan tap-out TransJakarta periode {stats['period']} dengan {fmt_int(stats['raw_rows'])} baris dan {stats['raw_cols']} atribut. Berdasarkan Tabel 4.1, atribut yang langsung berkaitan dengan analisis rute dan waktu meliputi corridorID, corridorName, tapInStopsName, tapOutStopsName, tapInTime, dan tapOutTime. Pemeriksaan kualitas data dilakukan untuk melihat missing value dan distribusi volume dasar perjalanan sebelum praproses lanjutan.",
    )
    add_figure(
        doc,
        1,
        "01_top_10_missing_value.png",
        "Top 10 missing value pada dataset awal",
        f"atribut dengan missing value tertinggi adalah {top_missing.iloc[0]['kolom']} sebanyak {fmt_int(top_missing.iloc[0]['missing_values'])} baris atau {fmt_pct(top_missing.iloc[0]['pct'])}. Posisi berikutnya adalah {top_missing.iloc[1]['kolom']} sebanyak {fmt_int(top_missing.iloc[1]['missing_values'])} baris dan {top_missing.iloc[2]['kolom']} sebanyak {fmt_int(top_missing.iloc[2]['missing_values'])} baris. Komposisi ini menunjukkan bahwa masalah kualitas data paling besar berada pada identitas koridor dan atribut tap-out.",
    )
    add_figure(
        doc,
        2,
        "02_top_10_koridor.png",
        "Top 10 koridor dengan frekuensi perjalanan tertinggi",
        f"koridor paling sering muncul adalah {top_corridors.index[0]} sebanyak {fmt_int(top_corridors.iloc[0])} perjalanan. Dua koridor setelahnya adalah {top_corridors.index[1]} sebanyak {fmt_int(top_corridors.iloc[1])} perjalanan dan {top_corridors.index[2]} sebanyak {fmt_int(top_corridors.iloc[2])} perjalanan. Konsentrasi volume pada beberapa koridor ini memberi indikasi awal bahwa permintaan perjalanan tidak tersebar merata di seluruh jaringan.",
    )
    add_figure(
        doc,
        3,
        "03_top_10_halte_asal.png",
        "Top 10 halte asal dengan frekuensi tap-in tertinggi",
        f"halte asal yang paling dominan adalah {top_tapin.index[0]} dengan {fmt_int(top_tapin.iloc[0])} transaksi tap-in. Peringkat kedua dan ketiga ditempati {top_tapin.index[1]} sebanyak {fmt_int(top_tapin.iloc[1])} transaksi dan {top_tapin.index[2]} sebanyak {fmt_int(top_tapin.iloc[2])} transaksi. Pola ini menandakan bahwa keberangkatan penumpang banyak terkonsentrasi pada simpul tertentu seperti Penjaringan, BKN, dan BNN LRT.",
    )
    add_figure(
        doc,
        4,
        "04_top_10_halte_tujuan.png",
        "Top 10 halte tujuan dengan frekuensi tap-out tertinggi",
        f"halte tujuan paling dominan adalah {top_tapout.index[0]} dengan {fmt_int(top_tapout.iloc[0])} transaksi tap-out. Peringkat berikutnya adalah {top_tapout.index[1]} sebanyak {fmt_int(top_tapout.iloc[1])} transaksi dan {top_tapout.index[2]} sebanyak {fmt_int(top_tapout.iloc[2])} transaksi. Perbedaan volume tap-in dan tap-out ini menunjukkan adanya simpul tujuan yang lebih kuat dibanding titik asal tertentu.",
    )

    add_text(doc, "4.1.2 Transformasi Variabel Waktu", bold=True, indent=False)
    add_text(
        doc,
        "Transformasi variabel waktu mengubah kolom waktu mentah menjadi tanggal perjalanan, jam tap-in, jam tap-out, dan durasi perjalanan. Hasil transformasi inilah yang menjadi dasar analisis temporal pada tahap berikutnya.",
    )
    add_figure(
        doc,
        5,
        "05_distribusi_jam_tap_in.png",
        "Distribusi jam tap-in pada hari kerja dan akhir pekan",
        f"puncak frekuensi tap-in secara keseluruhan berada pada pukul {hour_label(peak_hours.index[0])} sebanyak {fmt_int(peak_hours.iloc[0])} transaksi, diikuti pukul {hour_label(peak_hours.index[1])} sebanyak {fmt_int(peak_hours.iloc[1])} transaksi dan pukul {hour_label(peak_hours.index[2])} sebanyak {fmt_int(peak_hours.iloc[2])} transaksi. Pada hari kerja, puncak tertinggi terjadi pukul {hour_label(weekday_peak.index[0])} dengan {fmt_int(weekday_peak.iloc[0])} transaksi, sedangkan pada akhir pekan puncaknya bergeser ke pukul {hour_label(weekend_peak.index[0])} dengan {fmt_int(weekend_peak.iloc[0])} transaksi. Pola ini memperlihatkan kombinasi mobilitas pagi dan sore yang kuat pada hari kerja serta ritme yang lebih terkonsentrasi di pagi hari pada akhir pekan.",
    )

    add_text(doc, "4.1.3 Filter Durasi Perjalanan", bold=True, indent=False)
    add_text(
        doc,
        f"Filter durasi diterapkan untuk menyingkirkan observasi yang berada di luar rentang perjalanan wajar. Tahap ini menurunkan data dari {fmt_int(stage_counts['Setelah parsing waktu'])} menjadi {fmt_int(stage_counts['Setelah filter durasi'])} baris, sehingga {fmt_int(stats['removed_duration'])} baris dieliminasi.",
    )
    add_figure(
        doc,
        6,
        "06_distribusi_durasi_perjalanan.png",
        "Distribusi durasi perjalanan setelah transformasi waktu",
        f"rentang durasi yang paling dominan adalah 60-90 menit sebanyak {fmt_int(duration_table[2][1])} perjalanan atau {fmt_pct(duration_table[2][2])}, sangat dekat dengan rentang 30-60 menit sebanyak {fmt_int(duration_table[1][1])} perjalanan atau {fmt_pct(duration_table[1][2])}. Rentang 90-120 menit masih tinggi, yaitu {fmt_int(duration_table[3][1])} perjalanan atau {fmt_pct(duration_table[3][2])}, sedangkan perjalanan 120-180 menit hanya {fmt_int(duration_table[4][1])} perjalanan atau {fmt_pct(duration_table[4][2])}. Artinya, mayoritas perjalanan terkonsentrasi pada durasi menengah antara 30 sampai 120 menit.",
    )
    add_figure(
        doc,
        7,
        "07_filter_durasi_perjalanan.png",
        "Hasil filter durasi perjalanan",
        f"visual ini memperlihatkan bahwa data ekstrem di luar ambang durasi penelitian tidak dipertahankan. Setelah penyaringan, jumlah observasi valid menjadi {fmt_int(stage_counts['Setelah filter durasi'])} baris. Dengan demikian, distribusi durasi yang masuk ke model lebih merepresentasikan pola perjalanan reguler daripada outlier yang jarang terjadi.",
    )

    add_text(doc, "4.1.4 Filter Jam Operasional", bold=True, indent=False)
    add_text(
        doc,
        f"Filter jam operasional membatasi transaksi pada jam layanan TransJakarta yang relevan untuk penelitian. Tahap ini menurunkan data dari {fmt_int(stage_counts['Setelah filter durasi'])} menjadi {fmt_int(stage_counts['Setelah filter jam'])} baris atau menyaring {fmt_int(stats['removed_hour'])} baris tambahan.",
    )
    add_figure(
        doc,
        8,
        "08_filter_jam_operasional.png",
        "Hasil filter jam operasional",
        f"setelah pembatasan jam operasional, dataset menyisakan {fmt_int(stage_counts['Setelah filter jam'])} observasi. Tahap ini penting karena pola mobilitas yang dianalisis kemudian benar-benar merefleksikan perjalanan pada saat layanan aktif, bukan transaksi di luar rentang operasional.",
    )

    add_text(doc, "4.1.5 Feature Engineering", bold=True, indent=False)
    add_text(
        doc,
        "Feature engineering membentuk variabel turunan seperti day_of_week, is_weekend, n_trips, n_days_month, dan is_commuter. Variabel-variabel ini dipakai untuk membaca intensitas penggunaan dan kebiasaan perjalanan penumpang secara lebih rinci.",
    )
    add_figure(
        doc,
        9,
        "09_distribusi_intensitas_perjalanan.png",
        "Distribusi intensitas perjalanan per pengguna per hari",
        f"nilai n_trips yang paling dominan adalah 2 trip per hari dengan {fmt_int(intensity_table[1][1])} observasi pengguna-hari atau {fmt_pct(intensity_table[1][2])}. Nilai 1 trip per hari berada tepat di bawahnya sebanyak {fmt_int(intensity_table[0][1])} observasi atau {fmt_pct(intensity_table[0][2])}. Sementara itu, nilai 5 dan 6 trip hanya berjumlah {fmt_int(intensity_table[4][1])} dan {fmt_int(intensity_table[5][1])} observasi. Ini menunjukkan bahwa pola perjalanan harian mayoritas pengguna masih berada pada rentang 1 sampai 2 perjalanan.",
    )
    add_figure(
        doc,
        10,
        "10_feature_engineering_hari.png",
        "Hasil feature engineering pada dimensi hari",
        f"proporsi transaksi hari kerja pada dataset final mencapai {fmt_pct(stats['weekday_pct'])}, sedangkan transaksi akhir pekan sebesar {fmt_pct(stats['weekend_pct'])}. Selain itu, proporsi pengguna yang masuk kategori commuter mencapai {fmt_pct(stats['commuter_pct'])}. Temuan ini menegaskan bahwa data final didominasi perjalanan rutin hari kerja.",
    )

    add_text(doc, "4.1.6 Imputasi Nilai Hilang", bold=True, indent=False)
    add_text(
        doc,
        "Imputasi dilakukan untuk menutup missing value yang masih tersisa setelah feature engineering, terutama pada atribut koridor. Pendekatan yang digunakan memanfaatkan informasi halte yang sudah tersedia pada baris transaksi yang sama.",
    )
    add_figure(
        doc,
        11,
        "11_imputasi_nilai_hilang.png",
        "Hasil imputasi nilai hilang pada atribut koridor",
        f"jumlah missing value pada corridorName turun dari {fmt_int(stats['corridor_missing_before'])} baris menjadi {fmt_int(stats['corridor_missing_after'])} baris setelah imputasi, atau berkurang {fmt_pct(stats['imputation_reduction_pct'])}. Pada tahap ini missing value tapOutStopsName juga sudah turun menjadi {fmt_int(stats['tapout_missing_after'])} baris. Angka tersebut menunjukkan bahwa imputasi berhasil menutup hampir seluruh kekosongan informasi koridor sebelum proses cleaning akhir.",
    )

    add_text(doc, "4.1.7 Pembersihan Data dan Outlier", bold=True, indent=False)
    add_text(
        doc,
        f"Tahap cleaning dan penanganan outlier dilakukan untuk membuang sisa observasi yang belum konsisten. Jumlah data berkurang dari {fmt_int(stage_counts['Setelah imputasi'])} menjadi {fmt_int(stage_counts['Setelah cleaning'])} baris sehingga {fmt_int(stats['removed_cleaning'])} baris dihapus pada tahap ini.",
    )
    add_figure(
        doc,
        12,
        "12_data_cleaning_funnel.png",
        "Funnel data cleaning",
        f"penyusutan data berlangsung bertahap dari {fmt_int(stage_counts['Dataset awal'])} baris pada dataset awal menjadi {fmt_int(stage_counts['Dataset final'])} baris pada dataset akhir. Setelah cleaning, missing value corridorName sudah menjadi {fmt_int(stats['corridor_missing_final'])} baris. Artinya, dataset yang masuk ke tahap normalisasi sudah bebas dari kekosongan informasi utama pada atribut koridor.",
    )

    add_text(doc, "4.1.8 Normalisasi Variabel", bold=True, indent=False)
    add_text(
        doc,
        "Normalisasi z-score diterapkan pada variabel kontinu agar semua fitur berada pada skala yang sebanding. Langkah ini penting untuk mencegah fitur dengan rentang besar mendominasi proses clustering probabilistik.",
    )
    add_figure(
        doc,
        13,
        "13_normalisasi_variabel.png",
        "Normalisasi z-score pada variabel kontinu",
        f"normalisasi tidak mengubah jumlah observasi karena dataset tetap berisi {fmt_int(stage_counts['Setelah normalisasi'])} baris. Hasil standardisasi ini digunakan langsung sebagai feature matrix untuk GMM, khususnya pada variabel z_tapIn_hour, z_duration_minutes, z_n_trips, dan z_n_days_month.",
    )
    add_figure(
        doc,
        14,
        "14_heatmap_korelasi_fitur.png",
        "Heatmap korelasi antar fitur pada dataset final",
        f"korelasi tertinggi terdapat antara {corr_pairs[0][0]} dan {corr_pairs[0][1]} sebesar {fmt_num(corr_pairs[0][2], 4)}, diikuti {corr_pairs[1][0]} dan {corr_pairs[1][1]} sebesar {fmt_num(corr_pairs[1][2], 4)}. Korelasi antara tapIn_hour dan duration_minutes berada pada {fmt_num(stats['corr_tap_duration'], 4)}, sehingga hubungan keduanya bersifat sedang dan masih menyimpan informasi yang berbeda. Dengan demikian, kombinasi fitur waktu, durasi, intensitas, dan kebiasaan perjalanan tetap layak dipakai bersama pada tahap pemodelan.",
    )

    add_doc_table(
        doc,
        3,
        "Fitur yang Digunakan pada Pemodelan Gaussian Mixture Model",
        ["Fitur", "Jenis", "Keterangan"],
        stats["gmm_feature_table"],
    )

    add_text(doc, "4.2 Hasil Pemodelan Gaussian Mixture Model (GMM)", bold=True, indent=False)
    add_text(
        doc,
        f"Pemodelan GMM dilakukan pada {fmt_int(stats['clean_rows'])} observasi menggunakan enam fitur seperti pada Tabel 4.3. Tujuan tahap ini adalah membentuk segmen mobilitas penumpang yang bersifat probabilistik, bukan pengelompokan kaku satu kelas satu aturan.",
    )

    add_text(doc, "4.2.1 Representasi Data untuk Pemodelan", bold=True, indent=False)
    add_text(
        doc,
        "Representasi data untuk pemodelan menggabungkan empat fitur kontinu yang telah dinormalisasi dan dua fitur biner. Struktur ini membuat model dapat membaca pola waktu keberangkatan, durasi, intensitas perjalanan, kebiasaan aktif bulanan, kecenderungan akhir pekan, dan status commuter secara bersamaan.",
    )

    add_text(doc, "4.2.2 Estimasi Parameter Menggunakan Expectation-Maximization (EM)", bold=True, indent=False)
    add_text(
        doc,
        "Estimasi parameter dilakukan dengan algoritma Expectation-Maximization. Pada tahap ini setiap observasi diberi probabilitas keanggotaan untuk seluruh cluster, lalu parameter model diperbarui secara iteratif sampai konvergen. Pendekatan ini sesuai dengan karakter data mobilitas yang tidak selalu terpisah tegas.",
    )

    add_text(doc, "4.2.3 Seleksi Model dan Penentuan Jumlah Cluster Optimal", bold=True, indent=False)
    add_text(
        doc,
        f"File hasil seleksi model menetapkan konfigurasi akhir pada K={stats['selected_k']} dengan tipe kovarians {stats['selected_model']}. Penetapan ini kemudian dibandingkan dengan silhouette score agar jumlah cluster final tidak hanya mengacu pada satu indikator, tetapi juga pada kualitas pemisahan cluster.",
    )
    add_figure(
        doc,
        15,
        "15_kurva_bic_vs_jumlah_klaster.png",
        "Kurva BIC terhadap jumlah cluster",
        f"file seleksi model penelitian menetapkan K={stats['selected_k']} sebagai model akhir dengan nilai BIC {fmt_num(stats['selected_bic'])}. Pada tabel eksperimen BIC per K, nilai untuk K=4 adalah {fmt_num(stats['bic_map'][4])}, K=5 adalah {fmt_num(stats['bic_map'][5])}, dan K=6 adalah {fmt_num(stats['bic_map'][6])}. Oleh karena itu, keputusan model final dibaca bersama dengan indikator evaluasi lain agar segmentasi yang dipilih tetap interpretatif.",
    )
    add_figure(
        doc,
        16,
        "16_delta_bic_antar_klaster.png",
        "Perubahan BIC antar jumlah cluster",
        f"perubahan BIC dari K=4 ke K=5 sebesar {fmt_num(stats['bic_delta_45'])}, sedangkan perubahan dari K=5 ke K=6 sebesar {fmt_num(stats['bic_delta_56'])}. Nilai ini menunjukkan bahwa penambahan cluster setelah K=5 masih mengubah kecocokan model, namun perubahan tersebut tidak otomatis menghasilkan pemisahan cluster yang lebih baik ketika dibandingkan dengan metrik silhouette.",
    )
    add_figure(
        doc,
        17,
        "17_silhouette_score_per_jumlah_klaster.png",
        "Silhouette score per jumlah cluster",
        f"nilai silhouette tertinggi tercatat pada K=5, yaitu {fmt_num(stats['silhouette_k5'], 4)}. Nilai ini lebih tinggi dibanding K=4 sebesar {fmt_num(stats['silhouette_k4'], 4)} dan K=6 sebesar {fmt_num(stats['silhouette_k6'], 4)}. Karena itu, lima cluster dipertahankan sebagai kompromi terbaik antara kecocokan model dan keterpisahan cluster.",
    )

    add_text(doc, "4.2.4 Penugasan Cluster Probabilistik", bold=True, indent=False)
    add_text(
        doc,
        "Setelah jumlah cluster ditetapkan, model menghasilkan probabilitas posterior untuk setiap observasi. Distribusi probabilitas maksimum dipakai untuk menilai seberapa tegas model menempatkan observasi ke dalam cluster akhir.",
    )
    add_figure(
        doc,
        18,
        "18_distribusi_probabilitas_posterior.png",
        "Distribusi probabilitas posterior maksimum",
        f"median probabilitas posterior maksimum berada pada {fmt_num(stats['posterior_median'], 3)} dan kuartil ketiganya juga berada pada {fmt_num(stats['posterior_p75'], 3)}. Sebanyak {fmt_pct(stats['posterior_ge_90'])} observasi memiliki probabilitas maksimum minimal 0,90, sedangkan {fmt_pct(stats['posterior_ge_80'])} observasi memiliki probabilitas minimal 0,80. Ini menandakan bahwa assignment cluster yang dihasilkan model cukup tegas untuk mayoritas data.",
    )

    add_text(doc, "4.2.5 Evaluasi Kualitas Clustering", bold=True, indent=False)
    add_text(
        doc,
        "Evaluasi kualitas clustering tidak hanya berhenti pada silhouette score. Penelitian juga melihat keseimbangan ukuran cluster, entropy probabilitas, dan skor komposit untuk memastikan kualitas cluster dari beberapa sisi sekaligus.",
    )
    add_figure(
        doc,
        19,
        "19_evaluasi_kualitas_clustering.png",
        "Ringkasan evaluasi kualitas clustering",
        f"pada K=5, nilai cluster balance tercatat {fmt_num(stats['balance_k5'], 4)}, entropy sebesar {fmt_num(stats['entropy_k5'], 4)}, dan composite score sebesar {fmt_num(stats['composite_k5'], 4)}. Nilai entropy yang rendah menunjukkan ketidakpastian assignment yang kecil, sedangkan cluster balance yang moderat menunjukkan distribusi ukuran cluster masih cukup layak untuk diinterpretasikan.",
    )

    add_doc_table(
        doc,
        4,
        "Ringkasan Profil Cluster Hasil GMM",
        ["Label Cluster", "Jumlah", "Proporsi", "Jam Tap-In Rata-rata", "Durasi Rata-rata", "Weekend", "Commuter", "Koridor Dominan"],
        [
            [
                label,
                fmt_int(row["n_obs"]),
                fmt_pct(row["pct_obs"]),
                fmt_num(row["mean_tapIn_hour"]),
                fmt_num(row["mean_duration_min"]),
                fmt_pct(row["pct_weekend"]),
                fmt_pct(row["pct_commuter"]),
                row["top_corridor"],
            ]
            for label, row in cluster_profiles.iterrows()
        ],
    )

    add_text(doc, "4.2.6 Output Pemodelan dan Interpretasi Cluster", bold=True, indent=False)
    add_text(
        doc,
        "Interpretasi cluster dibaca dari ukuran cluster, rata-rata jam tap-in, rata-rata durasi, intensitas penggunaan, proporsi weekend, dan status commuter. Hasilnya menunjukkan bahwa model memisahkan penumpang rutin pagi, penumpang rutin sore, penumpang pagi dini, penumpang intensif, dan penumpang kasual.",
    )
    add_figure(
        doc,
        20,
        "20_distribusi_ukuran_klaster.png",
        "Distribusi ukuran cluster hasil GMM",
        f"cluster terbesar adalah {stats['largest_cluster']['label']} dengan {fmt_int(stats['largest_cluster']['n_obs'])} observasi atau {fmt_pct(stats['largest_cluster']['pct_obs'])}. Cluster kedua terbesar adalah Commuter Pagi Dini dengan {fmt_int(cluster_profiles.loc['Commuter Pagi Dini', 'n_obs'])} observasi atau {fmt_pct(cluster_profiles.loc['Commuter Pagi Dini', 'pct_obs'])}. Sementara itu, cluster terkecil adalah {stats['smallest_cluster']['label']} dengan {fmt_int(stats['smallest_cluster']['n_obs'])} observasi atau {fmt_pct(stats['smallest_cluster']['pct_obs'])}. Distribusi ini menunjukkan bahwa segmen sore dan pagi dini mendominasi data perjalanan.",
    )
    add_figure(
        doc,
        21,
        "21_heatmap_karakteristik_klaster.png",
        "Heatmap karakteristik fitur per cluster",
        f"heatmap memperlihatkan bahwa Commuter Sore memiliki rata-rata jam tap-in tertinggi sebesar {fmt_num(cluster_profiles.loc['Commuter Sore', 'mean_tapIn_hour'])}, sedangkan Commuter Pagi Dini paling awal pada {fmt_num(cluster_profiles.loc['Commuter Pagi Dini', 'mean_tapIn_hour'])}. Dari sisi durasi, Commuter Pagi memiliki rata-rata tertinggi sebesar {fmt_num(cluster_profiles.loc['Commuter Pagi', 'mean_duration_min'])} menit, sementara Commuter Pagi Dini terendah sebesar {fmt_num(cluster_profiles.loc['Commuter Pagi Dini', 'mean_duration_min'])} menit. Pola ini menunjukkan perbedaan temporal dan beban perjalanan yang cukup jelas antar cluster.",
    )
    add_figure(
        doc,
        22,
        "22_pola_jam_tap_in_per_klaster.png",
        "Pola jam tap-in per cluster",
        f"Commuter Pagi Dini terkonsentrasi pada rata-rata jam {fmt_num(cluster_profiles.loc['Commuter Pagi Dini', 'mean_tapIn_hour'])}, Commuter Pagi pada {fmt_num(cluster_profiles.loc['Commuter Pagi', 'mean_tapIn_hour'])}, dan Commuter Sore pada {fmt_num(cluster_profiles.loc['Commuter Sore', 'mean_tapIn_hour'])}. Penumpang Kasual dan Penumpang Intensif cenderung muncul pada siang hari, masing-masing pada {fmt_num(cluster_profiles.loc['Penumpang Kasual', 'mean_tapIn_hour'])} dan {fmt_num(cluster_profiles.loc['Penumpang Intensif', 'mean_tapIn_hour'])}. Perbedaan ini mendukung penamaan cluster yang diberikan.",
    )
    add_figure(
        doc,
        23,
        "23_profil_weekend_dan_commuter.png",
        "Profil weekend dan commuter per cluster",
        f"Penumpang Kasual memiliki proporsi weekend {fmt_pct(cluster_profiles.loc['Penumpang Kasual', 'pct_weekend'])} dan proporsi commuter hanya {fmt_pct(cluster_profiles.loc['Penumpang Kasual', 'pct_commuter'])}. Sebaliknya, Commuter Sore, Commuter Pagi, dan Commuter Pagi Dini masing-masing memiliki proporsi commuter mendekati atau sama dengan 100%. Dengan demikian, model berhasil memisahkan perjalanan rutin dan non-rutin secara tegas.",
    )

    add_doc_table(
        doc,
        5,
        "Ringkasan Representasi Data Transaksi untuk ARM",
        ["Metrik", "Nilai"],
        [
            ["Connected trip total", fmt_int(stats["connected_trip_total"])],
            ["Proporsi connected trip terhadap dataset final", fmt_pct(stats["connected_pct_of_final"])],
            ["Jumlah kandidat pasangan perjalanan", fmt_int(stats["pair_candidate_total"])],
            ["Proporsi perjalanan lintas koridor", fmt_pct(stats["cross_pct"])],
            ["Proporsi perjalanan pada koridor yang sama", fmt_pct(stats["self_pct"])],
        ],
    )
    add_doc_table(
        doc,
        6,
        "Ringkasan Evaluasi Association Rule Mining",
        ["Metrik", "Nilai"],
        [
            ["Jumlah aturan global terpilih", fmt_int(stats["rules_selected_global"])],
            ["Trip coverage by rules", fmt_pct(stats["trip_coverage_pct"])],
            ["Rata-rata support global", fmt_pct(stats["avg_support_global_pct"])],
            ["Rata-rata support lokal", fmt_pct(stats["avg_support_local_pct"])],
            ["Rata-rata lift lokal", fmt_num(stats["avg_lift_local"], 2)],
            ["Median lift lokal", fmt_num(stats["median_lift_local"], 2)],
            ["Rata-rata lift global", fmt_num(stats["avg_lift_global"], 2)],
        ],
    )

    add_text(doc, "4.3 Hasil Association Rule Mining (ARM)", bold=True, indent=False)
    add_text(
        doc,
        "Tahap ARM dilakukan setelah setiap observasi memiliki label cluster dari GMM. Tujuannya adalah mengekstraksi pola keterkaitan antar rute atau antar pasangan perjalanan yang paling konsisten muncul pada masing-masing segmen mobilitas.",
    )

    add_text(doc, "4.3.1 Representasi Data Transaksi", bold=True, indent=False)
    add_text(
        doc,
        f"Berdasarkan Tabel 4.5, tahap ARM menggunakan {fmt_int(stats['connected_trip_total'])} connected trip atau {fmt_pct(stats['connected_pct_of_final'])} dari dataset final. Dari total tersebut, {fmt_pct(stats['cross_pct'])} merupakan perjalanan lintas koridor dan {fmt_pct(stats['self_pct'])} merupakan perjalanan pada koridor yang sama. Komposisi ini menunjukkan bahwa data transaksi ARM tetap menangkap pola perjalanan langsung maupun perjalanan berulang pada koridor yang sama.",
    )

    add_text(doc, "4.3.2 Pembentukan Transactions Object", bold=True, indent=False)
    add_text(
        doc,
        f"Transactions object dibentuk dari pasangan lhs dan rhs pada setiap perjalanan terhubung. Setiap transaksi mewakili satu pasangan origin-destination yang sudah dikaitkan dengan label cluster. Dengan total {fmt_int(stats['connected_trip_total'])} transaksi dan {fmt_int(stats['pair_candidate_total'])} kandidat pasangan, struktur ini cukup kaya untuk mendukung ekstraksi frequent itemset dan association rules.",
    )

    add_text(doc, "4.3.3 Analisis Frekuensi Item", bold=True, indent=False)
    add_text(
        doc,
        "Analisis frekuensi item digunakan untuk melihat koridor yang paling sering muncul pada tiap cluster sebelum aturan asosiasi dibentuk. Tahap ini membantu mengidentifikasi item dasar yang dominan dan memengaruhi pembentukan rules.",
    )
    add_figure(
        doc,
        24,
        "24_frekuensi_item_dominan_per_klaster.png",
        "Frekuensi item dominan per cluster",
        f"item paling dominan pada seluruh cluster adalah Matraman Baru - Ancol, tetapi intensitasnya berbeda. Support tertinggi untuk item ini muncul pada Penumpang Intensif sebesar {fmt_pct(top_item_cluster.loc['Penumpang Intensif', 'support'] * 100)}, diikuti Commuter Sore sebesar {fmt_pct(top_item_cluster.loc['Commuter Sore', 'support'] * 100)} dan Penumpang Kasual sebesar {fmt_pct(top_item_cluster.loc['Penumpang Kasual', 'support'] * 100)}. Ini menunjukkan bahwa beberapa koridor utama tetap mendominasi lintas segmen, meskipun kekuatan dominasinya berbeda.",
    )

    add_text(doc, "4.3.4 Penerapan Algoritma FP-Growth", bold=True, indent=False)
    add_text(
        doc,
        f"Algoritma FP-Growth dijalankan dengan parameter penyaringan bertahap. Threshold awal yang tersimpan pada file konfigurasi meliputi strict support {fmt_num(stats['filter_params']['strict_support'], 2)}, strict confidence {fmt_num(stats['filter_params']['strict_confidence'], 2)}, dan strict lift {fmt_num(stats['filter_params']['strict_lift'], 2)}. Proses adaptif kemudian menyesuaikan min-count global menjadi {fmt_int(stats['filter_params']['adaptive_min_count_global'])} dan min-count per cluster menjadi {fmt_int(stats['filter_params']['adaptive_min_count_cluster'])} agar jumlah aturan yang dihasilkan tetap berada pada rentang yang dapat diinterpretasikan.",
    )

    add_text(doc, "4.3.5 Evaluasi Association Rules", bold=True, indent=False)
    add_text(
        doc,
        "Evaluasi association rules dilakukan dengan membaca jumlah aturan terpilih, support, confidence, lift, dan cakupan perjalanan yang dapat dijelaskan oleh aturan. Hasil evaluasi lengkapnya dirangkum pada Tabel 4.6 dan visual berikut.",
    )
    add_figure(
        doc,
        25,
        "25_jumlah_rules_per_klaster.png",
        "Jumlah aturan terpilih pada setiap cluster",
        f"cluster dengan jumlah aturan terbanyak adalah Commuter Sore sebanyak {fmt_int(rules_by_cluster['Commuter Sore'])} aturan, diikuti Commuter Pagi Dini sebanyak {fmt_int(rules_by_cluster['Commuter Pagi Dini'])} aturan. Commuter Pagi hanya menghasilkan {fmt_int(rules_by_cluster['Commuter Pagi'])} aturan, sedangkan Penumpang Intensif dan Penumpang Kasual tidak menghasilkan aturan terpilih. Hal ini menunjukkan bahwa pola asosiasi paling stabil justru muncul pada segmen commuter.",
    )
    add_figure(
        doc,
        26,
        "26_plot_support_confidence_lift.png",
        "Plot support, confidence, dan lift aturan asosiasi",
        f"secara global, penelitian menghasilkan {fmt_int(stats['rules_selected_global'])} aturan terpilih dengan rata-rata support global {fmt_pct(stats['avg_support_global_pct'])}, rata-rata lift lokal {fmt_num(stats['avg_lift_local'])}, dan median lift lokal {fmt_num(stats['median_lift_local'])}. Sebaran titik memperlihatkan bahwa aturan yang dipertahankan cenderung mempunyai confidence menengah hingga tinggi dengan lift di atas 1, sehingga relasi yang dipilih bukan sekadar kebetulan frekuensi.",
    )

    add_text(doc, "4.3.6 Analisis Pola Asosiasi Per Cluster", bold=True, indent=False)
    add_text(
        doc,
        "Setelah aturan terpilih diperoleh, analisis dilanjutkan pada pola origin-destination dominan dan hubungan jaringan antar rute. Langkah ini membantu membaca apakah tiap cluster cenderung menghasilkan perjalanan berulang pada koridor yang sama atau perpindahan lintas koridor yang spesifik.",
    )
    add_figure(
        doc,
        27,
        "27_top_od_patterns_per_klaster.png",
        "Top origin-destination patterns per cluster",
        f"pada Commuter Sore, pola origin-destination tertinggi adalah {top_od_cluster.loc['Commuter Sore', 'lhs']} menuju {top_od_cluster.loc['Commuter Sore', 'rhs']} sebanyak {fmt_int(top_od_cluster.loc['Commuter Sore', 'count'])} perjalanan. Pada Commuter Pagi Dini, pola tertinggi adalah {top_od_cluster.loc['Commuter Pagi Dini', 'lhs']} menuju {top_od_cluster.loc['Commuter Pagi Dini', 'rhs']} sebanyak {fmt_int(top_od_cluster.loc['Commuter Pagi Dini', 'count'])} perjalanan. Sementara itu, Penumpang Intensif paling banyak didominasi rute {top_od_cluster.loc['Penumpang Intensif', 'lhs']} menuju {top_od_cluster.loc['Penumpang Intensif', 'rhs']} sebanyak {fmt_int(top_od_cluster.loc['Penumpang Intensif', 'count'])} perjalanan.",
    )
    add_figure(
        doc,
        28,
        "28_jaringan_aturan_asosiasi.png",
        "Jaringan aturan asosiasi terpilih",
        f"aturan dengan lift lokal tertinggi berasal dari cluster {top_rule['cluster_label']} dengan pola {top_rule['lhs']} menuju {top_rule['rhs']}. Aturan ini memiliki confidence {fmt_pct(float(top_rule['confidence']) * 100)} dan lift lokal {fmt_num(top_rule['lift_local'])}. Jaringan pada gambar memperlihatkan bahwa sebagian aturan kuat masih berpusat pada koridor yang sama, namun terdapat juga relasi lintas rute yang menghubungkan simpul-simpul mobilitas tertentu.",
    )

    add_text(doc, "4.4 Analisis Integratif", bold=True, indent=False)
    add_text(
        doc,
        "Bagian ini mengintegrasikan hasil GMM dan ARM. Fokus utamanya adalah melihat bagaimana segmen mobilitas yang dibentuk secara probabilistik berkaitan dengan pola temporal serta dominansi rute yang ditemukan melalui Association Rule Mining.",
    )

    add_text(doc, "4.4.1 Profil Temporal Klaster", bold=True, indent=False)
    add_text(
        doc,
        "Profil temporal digunakan untuk melihat segmen mana yang mendominasi pada pagi, siang, sore, dan malam. Pembacaan ini penting karena cluster GMM memang dibentuk terutama dari pola waktu keberangkatan, durasi, dan intensitas perjalanan.",
    )
    add_figure(
        doc,
        29,
        "29_profil_temporal_klaster.png",
        "Profil temporal klaster berdasarkan kategori jam",
        f"pada kategori pagi, cluster yang paling dominan adalah {pagi['label']} dengan proporsi {fmt_pct(pagi['pct'])}. Pada siang hari dominasi berpindah ke {siang['label']} sebesar {fmt_pct(siang['pct'])}. Pada sore dan malam, cluster yang mendominasi adalah {sore['label']} sebesar {fmt_pct(sore['pct'])} dan {malam['label']} sebesar {fmt_pct(malam['pct'])}. Hasil ini konsisten dengan label cluster yang dihasilkan GMM, terutama untuk Commuter Pagi Dini dan Commuter Sore.",
    )

    add_text(doc, "4.4.2 Perbandingan Dominasi Rute Antar Klaster", bold=True, indent=False)
    add_text(
        doc,
        "Dominasi rute per cluster membantu melihat apakah setiap segmen mobilitas bertumpu pada koridor yang sama atau justru memiliki konsentrasi rute berbeda. Perbandingan ini dipakai untuk menjembatani hasil cluster GMM dengan pola transaksi pada ARM.",
    )
    add_figure(
        doc,
        30,
        "30_heatmap_dominansi_rute_per_klaster.png",
        "Heatmap dominansi rute pada setiap cluster",
        f"Commuter Sore paling banyak ditopang oleh rute {route_dominance[route_dominance['cluster_label'] == 'Commuter Sore'].iloc[0]['lhs']} dengan proporsi {fmt_pct(route_dominance[route_dominance['cluster_label'] == 'Commuter Sore'].iloc[0]['pct'])}. Commuter Pagi Dini juga didominasi oleh {route_dominance[route_dominance['cluster_label'] == 'Commuter Pagi Dini'].iloc[0]['lhs']} sebesar {fmt_pct(route_dominance[route_dominance['cluster_label'] == 'Commuter Pagi Dini'].iloc[0]['pct'])}. Sementara itu, Penumpang Intensif paling terkonsentrasi pada {route_dominance[route_dominance['cluster_label'] == 'Penumpang Intensif'].iloc[0]['lhs']} sebesar {fmt_pct(route_dominance[route_dominance['cluster_label'] == 'Penumpang Intensif'].iloc[0]['pct'])}. Perbedaan ini memperlihatkan bahwa tiap segmen tidak hanya berbeda pada waktu perjalanan, tetapi juga pada fokus rute yang digunakan.",
    )

    add_text(doc, "4.4.3 Diskusi Temuan dan Kontribusi", bold=True, indent=False)
    add_text(
        doc,
        f"Secara integratif, hasil GMM menunjukkan bahwa dataset final paling banyak dihuni oleh Commuter Sore ({fmt_pct(cluster_profiles.loc['Commuter Sore', 'pct_obs'])}) dan Commuter Pagi Dini ({fmt_pct(cluster_profiles.loc['Commuter Pagi Dini', 'pct_obs'])}). Hasil ARM kemudian memperlihatkan bahwa pola asosiasi yang stabil terutama muncul pada kedua cluster commuter tersebut, terlihat dari jumlah aturan Commuter Sore sebanyak {fmt_int(arm_eval_cluster.loc['Commuter Sore', 'rules_selected'])} dan Commuter Pagi Dini sebanyak {fmt_int(arm_eval_cluster.loc['Commuter Pagi Dini', 'rules_selected'])}. Temuan ini menunjukkan bahwa penumpang rutin tidak hanya memiliki waktu keberangkatan yang konsisten, tetapi juga rute yang lebih berulang dan lebih mudah dibaca melalui association rules. Dari sisi praktis, informasi ini dapat membantu identifikasi rute padat pada jam tertentu, evaluasi kebutuhan armada, serta pembacaan pola perjalanan berulang pada simpul-simpul yang dominan.",
    )

    add_text(doc, "4.4.4 Keterbatasan Penelitian dan Saran", bold=True, indent=False)
    add_text(
        doc,
        f"Penelitian ini masih memiliki beberapa keterbatasan. Pertama, periode data hanya mencakup {stats['period']} sehingga pola musiman atau variasi antar bulan belum tercakup. Kedua, data yang digunakan telah melalui anonimisasi sehingga karakteristik sosial pengguna tidak dapat dianalisis lebih jauh. Ketiga, association rules yang kuat lebih banyak muncul pada segmen commuter, sementara dua cluster lain tidak menghasilkan aturan terpilih karena frekuensi lokalnya lebih tersebar. Penelitian selanjutnya dapat memperluas periode data, menguji stabilitas cluster antar bulan, serta membandingkan hasil GMM dengan model clustering probabilistik lain untuk memastikan konsistensi segmentasi.",
    )

    doc.save(DOCX_PATH)


def main() -> None:
    ensure_dirs()
    cleanup_stale_visuals()
    setup_style()
    loaded = load_data()
    copy_visuals()
    build_visuals(loaded)
    stats = compute_stats(loaded)
    build_docx(stats)


if __name__ == "__main__":
    main()
